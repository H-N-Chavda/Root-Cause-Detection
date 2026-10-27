"""Build the combined progress deck.

Part A: the two week7 (Meet2-Midsem) explainer decks (PCMCI+, VAR-LiNGAM), recoloured to the new palette.
Part B: traditional PC results on plant data, read from the committed pca_pc_t2/outputs files.

Run from the repo root:  .venv/bin/python presentations/build_combined_deck.py
Output: presentations/week9/rcd_combined_deck.pptx
"""
import copy
import csv
import json
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OLD = ROOT / "presentations" / "week7"
RES = ROOT / "pca_pc_t2" / "outputs"
OUT = ROOT / "presentations" / "week9" / "rcd_combined_deck.pptx"

# ---------- palette ----------
IVORY, PANEL, INSET, BORDER = "F7F5EF", "FFFFFF", "F1EEE6", "DAD4C6"
FOREST, FOREST_MID, SAGE, SAGE_D = "174A3A", "2F7A5C", "8DAA91", "6A8E70"
CHAR, MUTED, GOLD, GOLD_D = "292D32", "676B70", "B79A5B", "9C7A35"
BRICK, PLUM = "A84A3E", "7B5577"
GOLD_TINT, GREEN_TINT, RED_TINT = "F5EEDD", "E8F0E9", "F6E7E3"

# old dark-theme colour -> new colour (used on every copied explainer slide)
COLOR_MAP = {
    "0F1218": IVORY, "1A1F2A": PANEL, "151A24": INSET,
    "2A3242": BORDER, "2E3646": "E3DED2", "3B465A": "C9C2B2", "4C586B": "A39D8F",
    "8B97A8": MUTED, "E9EDF3": CHAR,
    "2A2416": GOLD_TINT, "15241E": GREEN_TINT, "251D22": RED_TINT,
    "3DDC97": FOREST_MID,  # P sensor, kept / true links
    "F2B33D": GOLD_D,      # T sensor, "our pair"
    "9B8CFF": PLUM,        # F sensor
    "FF6B6B": BRICK,       # cut / rejected
    "6FC3DF": SAGE_D,      # secondary highlight
}
THEME_MAP = {  # theme slots: dk1, lt1, dk2, lt2, accent1..6
    "E9EDF3": CHAR, "0F1218": IVORY, "8B97A8": MUTED, "1A1F2A": PANEL,
    "3DDC97": FOREST_MID, "F2B33D": GOLD_D, "9B8CFF": PLUM, "FF6B6B": BRICK,
    "6FC3DF": SAGE_D, "2A3242": BORDER,
}

# new plain titles for the explainer slides (old title -> new title)
PCMCI_TITLES = [
    "The toy example: three reactor sensors",
    "The pair we will follow: T and P",
    "PCMCI+ in four pictures",
    "The basic test: partial correlation",
    "Part 1: build the parent list of P(t)",
    "Why the Part 1 list is not final",
    "Part 2: build the test set Z for our link",
    "A parent list moves back in time with its node",
    "Part 2 rounds: where our link is removed",
    "Orientation: give each link a direction",
    "The pair, from start to finish",
]
VAR_TITLES = [
    "Same reactor, one change: non-Gaussian noise",
    "Two kinds of link need different tools",
    "The model in two forms",
    "VAR-LiNGAM in four pictures",
    "Step 1: fit the VAR and keep the residuals",
    "Step 2a: how one direction is decided",
    "Step 2b: find the order, one sensor at a time",
    "Steps 3 and 4: correct the lagged weights, then prune",
    "The recovered graph matches the true graph",
]
# old cross-references to slide numbers that no longer hold
TEXT_FIXES = {
    "The pair we picked on the second slide is finished": "The pair we picked at the start is finished",
    "the equations on slide 1,": "the equations on the setup slide,",
    "the equations on slide one put in": "the equations on the setup slide put in",
}

FLARE_SHORT = {
    "(LRflare.pv+LREflare.pv+99fi4601.daca.pv)/1000": "LR flare sum",
    "509fi5201.daca.pv/1000": "509FI5201",
    "509mw5101.daca.pv": "509MW5101",
    "(44ft2505.pv*509mw5101.daca.pv/22.414)/1000": "44FT2505",
}

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def rgb(h):
    return RGBColor.from_string(h)


def I(x):
    return Inches(x)


# ---------- drawing helpers (same look as the explainer decks) ----------
def panel(slide, x, y, w, h, fill=PANEL, line=BORDER, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.04762 if min(w, h) > 1 else 0.12
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, runs, size=13, color=CHAR, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=None, spc=None, line_gap=None):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides-dict)."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if line_gap:
            p.space_before = Pt(line_gap) if k else Pt(0)
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
            if o.get("font", font):
                f.name = o.get("font", font)
            if o.get("spc", spc):
                r._r.get_or_add_rPr().set("spc", str(o.get("spc", spc)))
    return tb


def label(slide, x, y, w, t, color=MUTED):
    """Small spaced capitals, like 'THE THREE SENSORS' in the explainers."""
    return text(slide, x, y, w, 0.3, t.upper(), size=11.5, color=color, bold=True, spc=150)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def chip(slide, x, y, d, t, fill=PANEL, line=FOREST_MID, color=FOREST_MID, size=12):
    s = panel(slide, x, y, d, d, fill=fill, line=line, lw=1.6, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(color)
    return s


def callout(slide, y, head, body, x=0.68, w=11.97, h=0.42):
    """The gold 'OUR PAIR' bar used at the bottom of the explainer slides."""
    panel(slide, x, y, w, h, fill=GOLD_TINT, line=GOLD_D, lw=1.0)
    text(slide, x + 0.22, y, w - 0.4, h, [[(head + "   ", {"bold": True, "color": GOLD_D}), (body, {})]],
         size=12.5, anchor=MSO_ANCHOR.MIDDLE)


# ---------- deck-level helpers ----------
def recolor(el):
    for c in el.iter(f"{{{A_NS}}}srgbClr"):
        v = c.get("val", "").upper()
        if v in COLOR_MAP:
            c.set("val", COLOR_MAP[v])


def fix_text(el):
    for t in el.iter(f"{{{A_NS}}}t"):
        for old, new in TEXT_FIXES.items():
            if t.text and old in t.text:
                t.text = t.text.replace(old, new)
        if t.text and "amber" in t.text:  # the amber colour is now gold
            t.text = t.text.replace("amber", "gold")
    # a few circles were drawn as roundRect with adj > 50000; PowerPoint clamps it, other viewers draw stars
    for gd in el.iter(f"{{{A_NS}}}gd"):
        m = re.fullmatch(r"val (\d+)", gd.get("fmla", ""))
        if gd.get("name") == "adj" and m and int(m.group(1)) > 50000:
            gd.set("fmla", "val 50000")


def set_ph(slide, idx, t, size, color_scheme_tx, bold=False):
    """Write text into the title (100) or subtitle (101) placeholder, keeping the explainer style."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            if t is None:
                ph._element.getparent().remove(ph._element)
                return
            tf = ph.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = rgb(color_scheme_tx)
            return


def retitle(slide, new_title):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 100:
            runs = ph.text_frame.paragraphs[0].runs
            runs[0].text = new_title
            for r in runs[1:]:
                r._r.getparent().remove(r._r)


def main():
    prs = Presentation(OLD / "pcmci_plus_walkthrough.pptx")
    var = Presentation(OLD / "var_lingam_walkthrough.pptx")

    # theme: new colours and fonts (Georgia titles, Calibri body)
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    m = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
    scheme = m.group(0)
    for old, new in THEME_MAP.items():
        scheme = scheme.replace(f'val="{old}"', f'val="{new}"')
    xml = xml[: m.start()] + scheme + xml[m.end():]
    xml = xml.replace('typeface="Cambria"', 'typeface="Georgia"')
    theme._blob = xml.encode("utf-8")

    content = prs.slide_layouts[3]  # CONTENT: title, subtitle, footer, slide number
    footer_proto = None
    for sh in content.shapes:
        if sh.has_text_frame and sh.text_frame.text == "PCMCI+":
            footer_proto = copy.deepcopy(sh._element)
            sh.text_frame.paragraphs[0].runs[0].text = ""  # footer now written per slide
    num_proto = copy.deepcopy([s for s in prs.slides[0].shapes if s.name == "Slide Number Placeholder 0"][0]._element)

    def add_footer(slide, t):
        el = copy.deepcopy(footer_proto)
        slide.shapes._spTree.append(el)
        el.findall(f".//{{{A_NS}}}t")[0].text = t
        recolor(el)

    def new_slide(title, sub, foot, number=True):
        s = prs.slides.add_slide(content)
        set_ph(s, 100, title, 30, CHAR, bold=True)
        set_ph(s, 101, sub, 15, MUTED)
        add_footer(s, foot)
        if number:
            s.shapes._spTree.append(copy.deepcopy(num_proto))
        return s

    def divider(kicker, title, sub, note=None):
        s = prs.slides.add_slide(content)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(FOREST)
        set_ph(s, 100, None, 0, CHAR)
        set_ph(s, 101, None, 0, CHAR)
        panel(s, 0.68, 2.05, 0.9, 0.06, fill=GOLD, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(s, 0.68, 2.3, 11, 0.4, kicker.upper(), size=14, color=SAGE, bold=True, spc=200)
        text(s, 0.68, 2.75, 11.5, 1.0, title, size=40, color=IVORY, bold=True, font="Georgia")
        text(s, 0.68, 3.85, 10.5, 1.0, sub, size=18, color="DCE6DE")
        if note:
            text(s, 0.68, 5.3, 10.5, 0.8, note, size=15, color=GOLD, bold=True)
        return s

    # ----- Part A: recolour PCMCI+ slides in place, copy VAR-LiNGAM slides -----
    pcmci_slides = list(prs.slides)
    for s, t in zip(pcmci_slides, PCMCI_TITLES):
        recolor(s._element)
        fix_text(s._element)
        retitle(s, t)
        add_footer(s, "Part A  ·  PCMCI+")

    var_slides = []
    for src, t in zip(var.slides, VAR_TITLES):
        s = prs.slides.add_slide(content)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        for el in list(src.shapes._spTree)[2:]:
            s.shapes._spTree.append(copy.deepcopy(el))
        recolor(s._element)
        fix_text(s._element)
        retitle(s, t)
        add_footer(s, "Part A  ·  VAR-LiNGAM")
        if src.has_notes_slide:
            s.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
        var_slides.append(s)

    # ----- data for Part B -----
    s1 = json.loads((RES / "data/s1_summary.json").read_text())
    s5 = json.loads((RES / "data/s5_summary.json").read_text())
    s4 = json.loads((RES / "data/s4_summary.json").read_text())
    variance = list(csv.DictReader(open(RES / "data/pca_variance.csv")))
    shares = list(csv.DictReader(open(RES / "tables/pca_unit_share.csv")))
    coefs = list(csv.DictReader(open(RES / "t2/coefficients.csv")))
    nodes = list(csv.DictReader(open(RES / "tables/root_cause_nodes.csv")))
    thresholds = list(csv.DictReader(open(RES / "tables/flare_thresholds.csv")))
    sweep = list(csv.DictReader(open(RES / "tables/pc_alpha_sweep.csv")))
    events = list(csv.DictReader(open(RES / "data/flare_events.csv")))
    node = {r["node"]: r for r in nodes}

    def foot_b(s, src):
        add_footer(s, "Part B  ·  Plant data")
        text(s, 3.2, 6.96, 8.0, 0.28, "Source: pca_pc_t2/outputs/" + src, size=9, color=MUTED)

    # ----- opening -----
    title = prs.slides.add_slide(content)
    title.background.fill.solid()
    title.background.fill.fore_color.rgb = rgb(FOREST)
    set_ph(title, 100, None, 0, CHAR)
    set_ph(title, 101, None, 0, CHAR)
    panel(title, 0.68, 2.2, 0.9, 0.06, fill=GOLD, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(title, 0.68, 2.5, 11.5, 1.2, "Root cause detection for flaring", size=44, color=IVORY, bold=True, font="Georgia")
    text(title, 0.68, 3.65, 10.5, 1.0, "How two causal discovery methods work, and our first results on plant data",
         size=20, color="DCE6DE")
    text(title, 0.68, 5.6, 6, 0.4, "Progress review  ·  27 October 2026", size=14, color=SAGE)

    toc = new_slide("What this deck covers",
                    "Two separate parts. Part A explains methods on a toy example. Part B shows what we ran on plant data.",
                    "Root cause detection")
    for x, head, kick, sub, rows, accent in [
        (0.68, "Part A", "Algorithm explainers",
         "Studied on a toy example with a known answer. Not used on plant data yet.",
         [("PCMCI+, step by step", "5 – 15"), ("VAR-LiNGAM, step by step", "16 – 24"),
          ("The methods side by side", "25")], PLUM),
        (6.84, "Part B", "Traditional PC on plant data",
         "What we run today: PCA, then traditional PC, then T², on May 2025 data.",
         [("Our pipeline", "27"), ("Data and PCA", "28 – 29"), ("The causal graph", "30"),
          ("Problems found", "31 – 33"), ("Next steps", "34")], FOREST),
    ]:
        panel(toc, x, 1.75, 5.81, 4.75)
        panel(toc, x, 1.75, 5.81, 0.12, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(toc, x + 0.35, 2.1, 5.2, 0.35, head.upper(), size=12, color=accent, bold=True, spc=150)
        text(toc, x + 0.35, 2.45, 5.2, 0.5, kick, size=21, bold=True, font="Georgia")
        text(toc, x + 0.35, 3.05, 5.1, 0.6, sub, size=13, color=MUTED)
        for k, (r, n) in enumerate(rows):
            yy = 3.8 + k * 0.5
            panel(toc, x + 0.35, yy + 0.43, 5.1, 0.01, fill=BORDER, line=None, shape=MSO_SHAPE.RECTANGLE)
            text(toc, x + 0.35, yy + 0.06, 3.8, 0.35, r, size=14)
            text(toc, x + 4.15, yy + 0.06, 1.3, 0.35, "slide " + n, size=12, color=MUTED, align=PP_ALIGN.RIGHT)

    goal = new_slide("We want to find which process variables drive flaring",
                     "The plant records hundreds of tags every minute. We look for the ones that push the flare up.",
                     "Root cause detection")
    steps = [("PLANT DATA", "Hundreds of tags from 10 units, one reading per minute"),
             ("CAUSAL DISCOVERY", "Learn which variable affects which, as a graph"),
             ("ROOT CAUSE", "When a flare starts, find the variables that explain it")]
    for k, (h, b) in enumerate(steps):
        x = 0.68 + k * 4.12
        panel(goal, x, 1.8, 3.6, 1.55)
        label(goal, x + 0.3, 2.0, 3.0, h, color=FOREST_MID)
        text(goal, x + 0.3, 2.4, 3.0, 0.9, b, size=15)
        if k < 2:
            arrow(goal, x + 3.66, 2.57, x + 4.06, 2.57, color=MUTED, width=2)
    panel(goal, 0.68, 3.7, 5.81, 1.75)
    label(goal, 0.98, 3.9, 5.3, "Part A  ·  what we studied", color=PLUM)
    text(goal, 0.98, 4.3, 5.3, 1.1, ["Two advanced methods, PCMCI+ and VAR-LiNGAM.",
                                       "They handle time lags and same-minute effects.",
                                       "Shown on a toy reactor with three sensors."], size=14, line_gap=4)
    panel(goal, 6.84, 3.7, 5.81, 1.75)
    label(goal, 7.14, 3.9, 5.3, "Part B  ·  what we ran", color=FOREST_MID)
    text(goal, 7.14, 4.3, 5.3, 1.1, ["Traditional PC on one month of plant data.",
                                       "PCA first, then PC, then T² to find the cause.",
                                       "The pipeline runs, but the results need fixes."], size=14, line_gap=4)
    callout(goal, 5.8, "NOTE", "Part B uses traditional PC. PCMCI+ and VAR-LiNGAM are not applied to plant data yet.")

    div_a = divider("Part A  ·  Algorithm explainers", "How two advanced methods work",
                    "PCMCI+ and VAR-LiNGAM, on a toy reactor with three sensors. We made the data, so we know the "
                    "true answer and can check every step.",
                    "These methods are studied only. We do not use them on plant data yet.")

    # ----- comparison slide -----
    cmp_ = new_slide("PCMCI+ and VAR-LiNGAM side by side",
                     "Both recover the four true links of the toy reactor. Traditional PC, used in Part B, has neither of their extras.",
                     "Part A  ·  Comparison")
    cols = [("Traditional PC", "used in Part B", FOREST, True),
            ("PCMCI+", "Part A", PLUM, False), ("VAR-LiNGAM", "Part A", PLUM, False)]
    rows = [
        ("Time lags", ["No. Each row is treated as independent.", "Yes, up to τmax (2 in the toy).",
                       "Yes, through the VAR (τmax 1 in the toy)."]),
        ("Main tool", ["Conditional independence tests", "Conditional independence tests, in two parts",
                       "Regression, then an independence check on residuals"]),
        ("Same-minute direction", ["Collider rules. Some edges can stay undirected.", "Time order, then collider rules",
                                   "Time order, then the shape of the noise"]),
        ("Output", ["A graph, no weights", "A graph, no weights", "A graph with link weights"]),
        ("Needs", ["Enough data per test, no hidden common cause", "Same graph every minute (stationary)",
                   "Non-Gaussian noise, no hidden common cause"]),
    ]
    x0, w0, cw, y0, rh = 0.68, 2.3, 3.15, 1.8, 0.8
    for k, (name, tag, colr, used) in enumerate(cols):
        x = x0 + w0 + 0.12 + k * (cw + 0.12)
        panel(cmp_, x, y0, cw, 0.62 + rh * len(rows) + 0.1,
              fill=GREEN_TINT if used else PANEL, line=FOREST if used else BORDER, lw=1.6 if used else 1.0)
        text(cmp_, x + 0.25, y0 + 0.14, cw - 0.5, 0.3, name, size=17, bold=True, color=colr, font="Georgia")
        text(cmp_, x + 0.25, y0 + 0.5, cw - 0.5, 0.25, tag.upper(), size=10, bold=True, color=MUTED, spc=120)
        for r, (_, vals) in enumerate(rows):
            text(cmp_, x + 0.25, y0 + 0.78 + r * rh, cw - 0.45, rh - 0.1, vals[k], size=12.5)
    for r, (h, _) in enumerate(rows):
        label(cmp_, x0, y0 + 0.8 + r * rh, w0, h)

    # ----- Part B -----
    div_b = divider("Part B  ·  Running the algorithm on plant data", "Traditional PC on plant data",
                    "One month of plant data: May 2025, one reading per minute. PCA reduces the tags, traditional PC "
                    "builds the graph, and T² looks for the cause of alarms.",
                    "We use traditional PC here, not PCMCI+ or VAR-LiNGAM.")

    # 27 pipeline
    pipe = new_slide("Our pipeline: PCA, then PC, then T²",
                     "Five stages. All seven scripts passed their checks, and the full run took under 30 seconds.",
                     "")
    stages = [("1", "Load and clean", "298 tags in, 273 kept", "1-minute grid, 44,640 rows"),
              ("2", "PCA", "270 process tags", "→ 26 components (PCs)"),
              ("3", "Traditional PC", "26 PCs + 3 flare tags", "Rule: a flare tag causes nothing"),
              ("4", "Make a DAG", "Pick one direction", "for every edge"),
              ("5", "T² decomposition", "Alarm when T² is high", "then ask which node is to blame")]
    for k, (n, h, a, b) in enumerate(stages):
        x = 0.68 + k * 2.42
        used = n == "3"
        panel(pipe, x, 1.85, 2.1, 2.15, fill=GREEN_TINT if used else PANEL, line=FOREST if used else BORDER,
              lw=1.6 if used else 1.0)
        chip(pipe, x + 0.25, 2.05, 0.42, n, line=FOREST if used else FOREST_MID,
             color=FOREST if used else FOREST_MID)
        text(pipe, x + 0.25, 2.62, 1.75, 0.4, h, size=15, bold=True)
        text(pipe, x + 0.25, 3.05, 1.75, 0.9, [a, b], size=12, color=MUTED)
        if k < 4:
            arrow(pipe, x + 2.13, 2.92, x + 2.39, 2.92, color=MUTED, width=2)
    panel(pipe, 0.68, 4.3, 5.81, 1.85)
    label(pipe, 0.98, 4.48, 5.3, "How we run traditional PC", color=FOREST_MID)
    a01 = sweep[0]
    text(pipe, 0.98, 4.85, 5.3, 1.3, [
        "causal-learn PC, Fisher-Z test, alpha 0.01",
        "Depth up to 3, every 10th row (4,464 rows)",
        f"{int(a01['n_ci_tests']):,} tests, {float(a01['seconds']):.0f} seconds",
        "No time lags: each minute is treated as independent"], size=13, line_gap=3)
    panel(pipe, 6.84, 4.3, 5.81, 1.85)
    label(pipe, 7.14, 4.48, 5.3, "T² decomposition", color=FOREST_MID)
    text(pipe, 7.14, 4.85, 5.3, 1.3, [
        "Method from Li, Jin and Shi (2008)",
        f"Alarm when T² > {s5['t2_lim']:.1f}",
        f"A node is to blame when its term > {s5['z_lim']:.2f}",
        "Normal behaviour is learned from a reference period"], size=13, line_gap=3)
    foot_b(pipe, "SUMMARY.md, logs/")

    # 28 funnel
    fun = new_slide("From 298 tags to 29 model inputs",
                    "Cleaning removed 25 tags. PCA then squeezed 270 process tags into 26 components.",
                    "")
    drops = s1["dropped_s1_by_reason"]
    steps_f = [(298, "tags loaded", "294 process + 4 flare, from 10 units"),
               (273, "tags after cleaning", f"{s1['n_dropped_s1']} dropped (see right)"),
               (273, "into the model", "270 process tags go to PCA, 3 flare tags go straight to PC"),
               (29, "graph nodes", "26 PCs + 3 flare tags")]
    maxw = 6.2
    for k, (n, h, b) in enumerate(steps_f):
        y = 1.85 + k * 1.05
        w = max(1.3, maxw * n / 298)
        x = 0.68 + (maxw - w) / 2
        fill = FOREST if k == 3 else (SAGE if k else "C9D6CB")
        s = panel(fun, x, y, w, 0.8, fill=fill, line=None)
        text(fun, x, y, w, 0.8, "270 + 3" if k == 2 else str(n), size=24, bold=True, color=IVORY if k == 3 else CHAR,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font="Georgia")
        text(fun, 0.68, y + 0.82, maxw, 0.22, f"{h}  ·  {b}", size=11, color=MUTED, align=PP_ALIGN.CENTER)
    panel(fun, 7.3, 1.85, 5.35, 2.35)
    label(fun, 7.6, 2.03, 4.8, "Why 25 tags were dropped")
    reasons = [("Constant", drops["constant"]), ("Too much missing data", drops["missing"]),
               ("Almost constant", drops["near-constant"]), ("Constant in the reference", drops["reference"])]
    for k, (r, n) in enumerate(reasons):
        y = 2.45 + k * 0.42
        text(fun, 7.6, y, 2.6, 0.3, r, size=12.5)
        panel(fun, 10.25, y + 0.04, 1.9 * n / 10, 0.24, fill=SAGE if k else FOREST_MID, line=None,
              shape=MSO_SHAPE.RECTANGLE)
        text(fun, 10.3 + 1.9 * n / 10, y, 0.4, 0.3, str(n), size=12.5, bold=True)
    panel(fun, 7.3, 4.35, 5.35, 1.85)
    label(fun, 7.6, 4.53, 4.8, "What the data looks like")
    text(fun, 7.6, 4.9, 4.8, 1.3, [
        "1 to 31 May 2025, one row per minute: 44,640 rows",
        "Only 2 short gaps were filled; nothing else imputed",
        [("1 of 4 flare tags dropped: ", {}), ("44FT2505", {"bold": True}), (" is 64.5% missing", {})]],
        size=13, line_gap=3)
    foot_b(fun, "data/s1_summary.json, logs/s1_load_clean.log")

    # 29 PCA
    pca = new_slide("26 PCs keep 80% of the variance",
                    "The variance is spread thin: PC1 holds only 19.6%. The LOUP unit leads most of the first PCs.",
                    "")
    panel(pca, 0.68, 1.75, 5.9, 4.45)
    label(pca, 0.98, 1.92, 5.3, "Cumulative variance kept, %")
    cd = CategoryChartData()
    cd.categories = [r["pc"].replace("PC", "") for r in variance]
    cd.add_series("Cumulative", [round(float(r["cumulative"]) * 100, 1) for r in variance])
    gf = pca.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, I(0.85), I(2.95), I(5.6), I(2.9), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.size = Pt(11)
    ch.font.color.rgb = rgb(MUTED)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 100, 20
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E3DED2")
    va.format.line.fill.background()
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(BORDER)
    ca.tick_labels.font.size = Pt(9)
    ser = ch.plots[0].series[0]
    ser.format.line.color.rgb = rgb(FOREST)
    ser.format.line.width = Pt(2.25)
    ser.smooth = False
    ser.marker.size = 5
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(FOREST)
    ser.marker.format.line.color.rgb = rgb(FOREST)
    for k, (i, h) in enumerate([(0, "PC1 alone"), (9, "first 10 PCs"), (25, "all 26 PCs")]):
        text(pca, 0.98 + k * 1.8, 2.25, 1.7, 0.6,
             [[(f"{float(variance[i]['cumulative']) * 100:.1f}%",
                {"size": 22, "bold": True, "font": "Georgia", "color": GOLD_D if i == 25 else FOREST})], h],
             size=11, color=MUTED)
    text(pca, 0.98, 5.85, 5.3, 0.25, "component number", size=10, color=MUTED, align=PP_ALIGN.CENTER)

    units = ["LOUP", "ISOM", "PG", "DIU", "PDA", "SEU1", "SEU2", "SEU3", "ISOM_3", "DAK"]
    panel(pca, 6.84, 1.75, 5.81, 4.45)
    label(pca, 7.14, 1.92, 5.3, "Share of each PC's loadings by unit, %")
    gx, gy, cwid, chgt = 7.95, 2.6, 0.44, 0.31
    for j, u in enumerate(units):
        text(pca, gx + j * cwid - 0.05, gy - 0.3, cwid + 0.1, 0.25, u.replace("ISOM_3", "ISOM3"), size=8.5,
             color=MUTED, align=PP_ALIGN.CENTER, bold=True)
    for i, r in enumerate(shares[:10]):
        y = gy + i * chgt
        text(pca, 7.14, y + 0.04, 0.75, 0.25, r["pc"], size=11, bold=True)
        vals = [float(r[u]) for u in units]
        top = max(vals)
        for j, v in enumerate(vals):
            t = min(v / 0.45, 1.0)
            base, dark = (0xF7, 0xF5, 0xEF), (0x17, 0x4A, 0x3A)
            col = "".join(f"{round(b + (d - b) * t):02X}" for b, d in zip(base, dark))
            cell = panel(pca, gx + j * cwid, y, cwid - 0.03, chgt - 0.03, fill=col, line=None,
                         shape=MSO_SHAPE.RECTANGLE)
            text(pca, gx + j * cwid, y, cwid - 0.03, chgt - 0.03, f"{v * 100:.0f}", size=9,
                 bold=v == top, color=IVORY if t > 0.5 else CHAR, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(pca, 7.14, 5.8, 5.3, 0.35, "Darker = larger share. Bold = the leading unit of that PC.", size=11,
         color=MUTED)
    foot_b(pca, "data/pca_variance.csv, tables/pca_unit_share.csv")

    # 30 graph
    gr = new_slide("All 50 graph edges go from a PC to a flare tag",
                   "There is no edge between two PCs. So the graph is really three regressions, one per flare tag.",
                   "")
    panel(gr, 0.68, 1.7, 11.97, 3.25)
    pcs = [f"PC{i}" for i in range(1, 27)]
    flares = ["(LRflare.pv+LREflare.pv+99fi4601.daca.pv)/1000", "509fi5201.daca.pv/1000", "509mw5101.daca.pv"]
    px = {p: 0.98 + i * 0.437 for i, p in enumerate(pcs)}
    fx = {f: 2.6 + k * 4.0 for k, f in enumerate(flares)}
    d = 0.36
    strongest = max(coefs, key=lambda r: abs(float(r["beta"])))
    for r in sorted(coefs, key=lambda r: abs(float(r["beta"]))):
        b = abs(float(r["beta"]))
        is_top = r is strongest
        arrow(gr, px[r["parent"]] + d / 2, 2.05 + d, fx[r["node"]] + 0.9, 4.05,
              color=GOLD if is_top else SAGE, width=0.5 + 5 * b, head=False)
    for p in pcs:
        chip(gr, px[p], 2.05, d, p.replace("PC", ""), size=9.5, line=FOREST_MID, color=FOREST_MID)
    text(gr, 0.98, 1.8, 3, 0.22, "26 PCs  (all are roots: nothing points into them)", size=10.5, color=MUTED)
    for f in flares:
        s = panel(gr, fx[f], 4.05, 1.8, 0.5, fill=FOREST, line=None)
        npar = sum(1 for r in coefs if r["node"] == f)
        text(gr, fx[f], 4.05, 1.8, 0.5, [[(FLARE_SHORT[f], {"bold": True}), (f"   {npar} parents", {"size": 10.5})]],
             size=13, color=IVORY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(gr, 0.98, 4.62, 11.4, 0.25,
         f"Line width = size of the effect. Strongest edge, in gold: {strongest['parent']} → "
         f"{FLARE_SHORT[strongest['node']]} (weight {float(strongest['beta']):.2f}).", size=10.5, color=MUTED)
    cards = [("Why no PC-PC edges", "PCA makes components uncorrelated by design. So the PC test finds no "
                                    "link between them. This is expected, not a bug."),
             ("What this means", "The graph cannot show a chain like A → B → flare. It only says which PCs "
                                 "move with each flare tag."),
             ("Stable with alpha", f"With alpha 0.05 we get {sweep[1]['n_edges']} edges, again all PC → flare. "
                                   f"No flare tag points to anything.")]
    for k, (h, b) in enumerate(cards):
        x = 0.68 + k * 4.04
        panel(gr, x, 5.1, 3.89, 1.15)
        label(gr, x + 0.25, 5.22, 3.4, h, color=FOREST_MID)
        text(gr, x + 0.25, 5.52, 3.45, 0.7, b, size=12)
    foot_b(gr, "graph/dag_edges.csv, t2/coefficients.csv, tables/pc_alpha_sweep.csv")

    # 31 flare rule
    fl = new_slide("The flare rule turns the whole month into one event",
                   "A minute is flagged when a flare tag is above median + 3 × MAD. Flags less than 5 minutes apart are joined.",
                   "")
    ev = events[0]
    panel(fl, 0.68, 1.75, 11.97, 2.75)
    tx0, tw = 2.75, 9.55
    label(fl, 0.98, 1.92, 6, "May 2025, one bar per step")
    for day in (1, 8, 15, 22, 29):
        xx = tx0 + tw * (day - 1) / 31
        text(fl, xx - 0.3, 2.22, 0.6, 0.22, f"{day} May", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
        panel(fl, xx, 2.45, 0.01, 1.85, fill="E3DED2", line=None, shape=MSO_SHAPE.RECTANGLE)
    flagged = 15.9
    text(fl, 0.98, 2.6, 1.7, 0.5, [[("Flagged minutes", {"bold": True})], "before joining"], size=12)
    panel(fl, tx0, 2.62, tw, 0.42, fill=INSET, line=None, shape=MSO_SHAPE.RECTANGLE)
    panel(fl, tx0, 2.62, tw * flagged / 100, 0.42, fill=SAGE, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(fl, tx0 + tw * flagged / 100 + 0.12, 2.62, 6, 0.42,
         f"{flagged}% of all minutes (total length, not their real positions)", size=12, anchor=MSO_ANCHOR.MIDDLE)
    text(fl, 0.98, 3.4, 1.7, 0.5, [[("Flare events", {"bold": True})], "after joining"], size=12)
    panel(fl, tx0, 3.42, tw, 0.42, fill=INSET, line=None, shape=MSO_SHAPE.RECTANGLE)
    st, en = (1 + 89 / 1440), (31 + 20.5 / 24)
    panel(fl, tx0 + tw * (st - 1) / 31, 3.42, tw * (en - st) / 31, 0.42, fill=BRICK, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    text(fl, tx0 + 0.4, 3.42, 8.5, 0.42,
         f"1 event: {int(ev['start'][8:10])} May {ev['start'][11:16]} to {int(ev['end'][8:10])} May {ev['end'][11:16]}  ·  "
         f"{s1['pct_flare_rows']:.1f}% of all rows", size=12.5, bold=True, color=IVORY, anchor=MSO_ANCHOR.MIDDLE)
    nonflare = round(s1["rows"] * (1 - s1["pct_flare_rows"] / 100))
    text(fl, 0.98, 4.05, 11.3, 0.3, f"Only about {nonflare} minutes in the whole month count as 'no flare'.",
         size=12, color=MUTED)

    panel(fl, 0.68, 4.65, 5.81, 1.6)
    label(fl, 0.98, 4.8, 5.3, "Share of minutes above the threshold")
    for k, r in enumerate(thresholds):
        y = 5.15 + k * 0.26
        nm = FLARE_SHORT[r["flare_tag"]] + ("  (dropped later)" if "44ft" in r["flare_tag"] else "")
        text(fl, 0.98, y, 3.3, 0.25, nm, size=11.5)
        text(fl, 4.4, y, 1.8, 0.25, f"{float(r['pct_above']):.2f}%", size=11.5, bold=True, align=PP_ALIGN.RIGHT)
    panel(fl, 6.84, 4.65, 5.81, 1.6, fill=RED_TINT, line=BRICK)
    label(fl, 7.14, 4.8, 5.3, "Why this matters", color=BRICK)
    text(fl, 7.14, 5.12, 5.3, 1.1, [
        "There is no flare-free period to learn normal behaviour from.",
        "So the reference became the first 60% of May (to 19 May 14:23), with flares inside it."],
        size=12.5, line_gap=3)
    foot_b(fl, "data/flare_events.csv, tables/flare_thresholds.csv, data/deviations.json")

    # 32 T² alarms
    t2 = new_slide("The T² alarms do not behave as expected",
                   "Alarms are far too common in the reference period, and the decomposition does not add up when it matters.",
                   "")
    panel(t2, 0.68, 1.75, 6.6, 4.5)
    label(t2, 0.98, 1.92, 6, "Share of rows with a T² alarm, %")
    # columns drawn as shapes, like the other bar visuals in this deck
    cats = ["Reference rows", "All rows", "Flare rows", f"No-flare rows (~{nonflare})"]
    vals = [s5["pct_ooc_reference"], s5["pct_ooc_all"], s5["pct_ooc_flare"], s5["pct_ooc_nonflare"]]
    base_y, full_h, col_w = 5.15, 2.55, 1.05
    panel(t2, 0.98, base_y, 6.0, 0.015, fill=BORDER, line=None, shape=MSO_SHAPE.RECTANGLE)
    exp_y = base_y - full_h * 5 / 100
    for xx in [1.1 + 0.2 * j for j in range(29)]:  # dotted 5% line
        panel(t2, xx, exp_y, 0.1, 0.015, fill=GOLD_D, line=None, shape=MSO_SHAPE.RECTANGLE)
    for k, (c, v, col) in enumerate(zip(cats, vals, [GOLD, SAGE, SAGE, BRICK])):
        x = 1.25 + k * 1.45
        h = full_h * v / 100
        panel(t2, x, base_y - h, col_w, h, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(t2, x - 0.2, base_y - h - 0.36, col_w + 0.4, 0.32, f"{v:.1f}%", size=15, bold=True,
             align=PP_ALIGN.CENTER, font="Georgia")
        text(t2, x - 0.2, base_y + 0.06, col_w + 0.4, 0.45, c, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    text(t2, 0.98, 5.7, 6.1, 0.5, [[("Dashed line: reference rows should alarm about 5% of the time. ", {"bold": True}),
                                     ("No-flare rows alarm more than flare rows.", {})]], size=12)
    panel(t2, 7.5, 1.75, 5.15, 2.6)
    label(t2, 7.8, 1.92, 4.6, "Does the decomposition add up?")
    text(t2, 7.8, 2.25, 4.6, 0.5, "The node terms should sum to the total T². We measure the gap (0 = perfect).",
         size=12, color=MUTED)
    for k, (h, v, c) in enumerate([("reference rows", s5["discrepancy"]["reference"]["median"], FOREST_MID),
                                   ("alarm rows", s5["discrepancy"]["T2_ooc"]["median"], BRICK)]):
        x = 7.8 + k * 2.35
        text(t2, x, 2.85, 2.2, 0.75, f"{v:.2f}", size=40, bold=True, color=c, font="Georgia")
        text(t2, x, 3.62, 2.2, 0.5, ["median gap,", h], size=12, color=MUTED)
    panel(t2, 7.5, 4.5, 5.15, 1.75)
    label(t2, 7.8, 4.67, 4.6, "The normal-period terms are not normal")
    text(t2, 7.8, 5.0, 4.6, 1.2, [
        f"{s5['ref_pct_beyond_limit_mean']:.2f}% of reference values pass the limit; {s5['ref_expected_pct']:.2f}% expected",
        "Flare tags have very heavy tails (excess kurtosis 50 and 15)",
        "Likely cause: the reference period has flares in it (not yet checked)"], size=12, line_gap=3)
    foot_b(t2, "data/s5_summary.json, tables/reference_check.csv")

    # 33 T² points to
    tp = new_slide("What the T² points to, and why we cannot trust it yet",
                   "We count how often each node is to blame during a T² alarm. These are leads, not answers.",
                   "")
    panel(tp, 0.68, 1.75, 6.6, 4.5)
    label(tp, 0.98, 1.92, 6, "Node is to blame, % of alarm rows")
    top = sorted(nodes, key=lambda r: -float(r["pct_significant_when_T2_ooc"]))[:8]
    names = [FLARE_SHORT.get(r["node"], r["node"]) + (f"  ({r['unit'][4:-1]})" if r["node"].startswith("PC") else "  (flare)")
             for r in top]
    for k, (nm, r) in enumerate(zip(names, top)):
        y = 2.35 + k * 0.46
        v = float(r["pct_significant_when_T2_ooc"])
        is_pc = r["node"].startswith("PC")
        text(tp, 0.98, y + 0.04, 2.2, 0.3, nm, size=12, bold=k == 0)
        panel(tp, 3.2, y + 0.05, 3.3 * v / 20, 0.28, fill=FOREST_MID if is_pc else SAGE, line=None,
              shape=MSO_SHAPE.RECTANGLE)
        text(tp, 3.28 + 3.3 * v / 20, y + 0.04, 0.7, 0.3, f"{v:.1f}%", size=12, bold=True)
    text(tp, 0.98, 6.0, 6.0, 0.22, [[("■ ", {"color": FOREST_MID}), ("PC (main unit)    ", {}),
                                      ("■ ", {"color": SAGE}), ("flare tag", {})]], size=10.5, color=MUTED)
    pc7, pc14 = node["PC7"], node["PC14"]
    cards = [
        ("PC7 fires in normal times too",
         f"PC7 is to blame in {float(pc7['pct_significant_reference']):.1f}% of reference rows, but only "
         f"{float(pc7['pct_significant_when_T2_ooc']):.1f}% of alarm rows. It is noise, not a flare signal.", RED_TINT, BRICK),
        ("The 'first alarm' is the first minute",
         "The single event's first alarm is at 1 May 00:00, the start of the data. So PC13 and PC7, flagged there, "
         "do not mark a flare start.", RED_TINT, BRICK),
        ("One lead worth checking",
         "PC14 loads on SEU3 pressure tags 42PC4215 and 42PC4501 (0.31 each). PC1 is spread over many units "
         "(about 0.11 per tag), so it names no single tag.", GOLD_TINT, GOLD_D),
    ]
    for k, (h, b, fillc, linec) in enumerate(cards):
        y = 1.75 + k * 1.52
        panel(tp, 7.5, y, 5.15, 1.4, fill=fillc, line=linec)
        label(tp, 7.75, y + 0.13, 4.7, h, color=linec)
        text(tp, 7.75, y + 0.45, 4.7, 0.95, b, size=12)
    foot_b(tp, "tables/root_cause_nodes.csv, tables/root_cause_by_event.csv")

    # 34 next steps
    nx = new_slide("Next steps: fix the data setup, then rerun",
                   "The pipeline works. The inputs to it need fixing before the root causes can be read.",
                   "")
    tracks = [
        ("pca_pc_t2  ·  PCA + traditional PC + T²", FOREST, [
            "Fix the flare rule so that flares become separate events",
            "Find a flare-free reference period to learn normal behaviour",
            "Check the unit folder names in .env (3 units were not found)",
            "Rerun. Check: reference alarms near 5%, decomposition gap small",
            "Then read the root causes, starting with PC14 and the SEU3 tags"]),
        ("causal_t2  ·  traditional PC on all 267 tags", PLUM, [
            "The PC stage crashed: an IndexError after 827 s, then a segfault",
            "It has about 13 million tests even at depth 1",
            "Run the three checks in command.md: traceback, one-thread run, versions",
            "Pick the fix from the result: threading, memory, or library version"]),
    ]
    for k, (h, colr, items) in enumerate(tracks):
        x = 0.68 + k * 6.16
        panel(nx, x, 1.75, 5.81, 4.5)
        panel(nx, x, 1.75, 5.81, 0.12, fill=colr, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(nx, x + 0.3, 2.05, 5.3, 0.35, h, size=15, bold=True, color=colr)
        for i, it in enumerate(items):
            y = 2.6 + i * 0.68
            chip(nx, x + 0.3, y, 0.36, str(i + 1), line=colr, color=colr, size=11)
            text(nx, x + 0.85, y + 0.02, 4.75, 0.62, it, size=13)
    foot_b(nx, "SUMMARY.md, causal_t2/outputs/LOG.md, command.md")

    # ----- final order -----
    order = [title, toc, goal, div_a] + pcmci_slides + var_slides + [cmp_, div_b, pipe, fun, pca, gr, fl, t2, tp, nx]
    ids = prs.slides._sldIdLst
    elems = {prs.part.related_part(sid.rId): sid for sid in ids}
    for sid in list(ids):
        ids.remove(sid)
    for s in order:
        ids.append(elems[s.part])
    assert len(ids) == 34, len(ids)
    for n, s in enumerate(order, start=1):  # cached slide-number text, for viewers that do not refresh fields
        for fld in s._element.iter(f"{{{A_NS}}}fld"):
            if fld.get("type") == "slidenum":
                fld.find(f"{{{A_NS}}}t").text = str(n)
    prs.save(OUT)
    print("saved", OUT, len(ids), "slides")


if __name__ == "__main__":
    main()
