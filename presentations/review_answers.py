"""Content and slides for Review_Answers.pptx.

One slide per question raised in the review, plus a title and a closing ask.
The deck is deliberately short: it exists to be explained out loud in a few
minutes, not read.

Every figure here traces to FEEDBACK_RESPONSE.md at the repository root, which
carries the full argument and the citations.
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

import theme as T

ACT = "Review follow-up"

TITLE = "Five Questions, Answered"
SUBTITLE = "Follow-up to the last review · causal discovery for root-cause detection"
TITLE_CHIPS = ["65 sources reviewed", "3 datasets", "August 2026"]


# ---------------------------------------------------------------------------
# Shared devices
# ---------------------------------------------------------------------------

def _question(s, number, question, answer, accent=T.TEAL):
    """Every answer slide opens the same way: the question, then the answer in
    one sentence. The reader should be able to stop after the band."""
    T.header(s, ACT, number, accent=accent)
    T.text(s, T.MARGIN_IN, T.MARGIN_IN + 0.62, T.col_w(11), 0.5, question,
           size=T.SZ_TITLE, color=T.INK, bold=True)

    y = T.MARGIN_IN + 1.30
    T.rect(s, T.MARGIN_IN, y, T.col_w(12), 0.62, fill=_soft(accent),
           outline=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.rect(s, T.MARGIN_IN, y, 0.045, 0.62, fill=accent)
    T.text(s, T.MARGIN_IN + 0.24, y + 0.155, T.col_w(12) - 0.45, 0.32, answer,
           size=T.SZ_SUB, color=T.INK, bold=True)
    return y + 0.92


def _soft(accent):
    return {T.TEAL: T.TEAL_SOFT, T.SLATE: T.SLATE_SOFT, T.AMBER: T.AMBER_SOFT}[accent]


def _colhead(s, y, cells, widths, xs):
    for cell, w, x in zip(cells, widths, xs):
        T.text(s, x, y, w, 0.24, cell, size=T.SZ_MICRO, color=T.MUTED,
               bold=True, space=1.2)
    T.line(s, T.MARGIN_IN, y + 0.28, T.SLIDE_W_IN - T.MARGIN_IN, y + 0.28,
           color=T.INK, pt=0.9)
    return y + 0.42


def _footnote(s, body):
    T.caption(s, T.MARGIN_IN, T.SLIDE_H_IN - T.MARGIN_IN - 0.42, T.col_w(12), body)


# ---------------------------------------------------------------------------
# 1 — Title
# ---------------------------------------------------------------------------

def slide_01_title(prs):
    s = T.blank(prs, dark=True)
    T.rect(s, 0.95, 2.30, 0.042, 1.90, fill=T.TEAL)
    T.text(s, 1.22, 2.26, T.col_w(9), 1.10, TITLE, size=46,
           color=T.ON_DARK, bold=True, spacing=0.95)
    T.text(s, 1.22, 3.58, T.col_w(8), 0.62, SUBTITLE, size=17,
           color=T.RGBColor(0x9A, 0xA2, 0xAA), spacing=1.2)
    T.line(s, 1.22, 4.24, 4.6, 4.24, color=T.RGBColor(0x36, 0x3D, 0x45), pt=0.75)

    x = T.SLIDE_W_IN - T.MARGIN_IN
    for label in reversed(TITLE_CHIPS):
        w = 0.115 * len(label) + 0.34
        T.pill(s, x - w, 5.85, w, 0.34, label, fill=None,
               outline=T.RGBColor(0x3A, 0x44, 0x4C),
               color=T.RGBColor(0x8E, 0x97, 0xA0), size=T.SZ_CAPTION, bold=False)
        x -= w + 0.16
    return s


# ---------------------------------------------------------------------------
# 2 — Q1  C++ implementations
# ---------------------------------------------------------------------------

CPP = [
    ("cuPC", "CUDA / C++", "500–1300× over serial CPU", "NVIDIA GPU"),
    ("GPUCSL", "CUDA, Python API", "9.5× vs pcalg, 19.8× vs bnlearn", "NVIDIA GPU"),
    ("PEPC", "C++, multicore", "134× vs Tetrad", "CPU only"),
    ("pcalg", "R with C++ core", "reference baseline", "R runtime"),
]


def slide_02_cpp(prs):
    s = T.blank(prs)
    y = _question(
        s, 2,
        "Is there a C++ implementation of PC?",
        "Yes — four established ones. But they change speed, not answers.",
        accent=T.SLATE)

    xs = [T.MARGIN_IN, T.col_x(2), T.col_x(5), T.col_x(9)]
    ws = [T.col_w(2), T.col_w(3), T.col_w(4), T.col_w(3)]
    y = _colhead(s, y, ["IMPLEMENTATION", "LANGUAGE", "REPORTED SPEED-UP", "NEEDS"], ws, xs)
    for name, lang, speed, needs in CPP:
        for cell, w, x, bold, col in zip(
                [name, lang, speed, needs], ws, xs,
                [True, False, False, False],
                [T.INK, T.MUTED, T.INK, T.MUTED]):
            T.text(s, x, y, w, 0.3, cell, size=T.SZ_BODY, color=col, bold=bold)
        T.line(s, T.MARGIN_IN, y + 0.36, T.SLIDE_W_IN - T.MARGIN_IN, y + 0.36,
               color=T.RULE, pt=0.5)
        y += 0.46

    y += 0.26
    T.text(s, T.MARGIN_IN, y, T.col_w(7), 0.3,
           "Why we are not adopting one", size=T.SZ_SMALL, color=T.SLATE, bold=True)
    T.text(s, T.MARGIN_IN, y + 0.32, T.col_w(7), 0.82,
           ["cuPC runs “an order-independent version of PC” — the same "
            "statistics we already have. These are scheduling wins, not "
            "statistical ones.",
            "They optimise runs that take hours. Ours takes 0.13 s."],
           size=T.SZ_BODY, color=T.INK, spacing=1.20)

    bx = T.col_x(8)
    T.rect(s, bx, y - 0.08, T.col_w(4), 1.16, fill=T.PAPER_DIM, outline=None,
           radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, bx + 0.20, y + 0.06, T.col_w(4) - 0.40, 0.24,
           "OUR RUNTIME TODAY", size=T.SZ_MICRO, color=T.MUTED, bold=True, space=1.2)
    for i, (label, val) in enumerate([("Tennessee Eastman", "0.13 s"),
                                      ("Ultra-processed food", "0.29 s")]):
        T.text(s, bx + 0.20, y + 0.40 + i * 0.32, T.col_w(2) + 0.2, 0.26, label,
               size=T.SZ_SMALL, color=T.INK)
        T.text(s, bx + T.col_w(4) - 1.00, y + 0.38 + i * 0.32, 0.80, 0.26, val,
               size=T.SZ_SUB, color=T.SLATE, bold=True, align=PP_ALIGN.RIGHT,
               font=T.MONO)

    _footnote(s, "Revisit at scale: 31 → 150 variables took independence tests "
                 "from 5,900 to 1.97 million.")
    return s


# ---------------------------------------------------------------------------
# 3 — Q2  the library stack
# ---------------------------------------------------------------------------

def slide_03_libraries(prs):
    s = T.blank(prs)
    y = _question(
        s, 3,
        "Which libraries should we build on?",
        "causal-learn and tigramite now — and yes, C++ can run under Python: GPUCSL.",
        accent=T.TEAL)

    tiers = [
        (T.TEAL, "USE NOW — correctness",
         [("causal-learn", "Pure Python, from CMU. The library CIPCaD-Bench itself "
                           "used for its published baseline, so our numbers become "
                           "comparable."),
          ("tigramite", "PCMCI+ and LPCMCI — the time-series methods in question 3.")]),
        (T.SLATE, "ADD AT SCALE — speed",
         [("GPUCSL", "Python API over CUDA/C++ kernels. C++ performance without "
                     "leaving Python.")]),
        (T.MUTED, "NOT VIABLE HERE",
         [("pcalg / bnlearn", "R only; no R on this machine."),
          ("Tetrad", "Needs JDK 21, we have Java 8."),
          ("cuPC", "Raw CUDA, no Python binding.")]),
    ]

    for accent, heading, rows in tiers:
        T.rect(s, T.MARGIN_IN, y + 0.02, 0.045, 0.24 + 0.40 * len(rows), fill=accent)
        T.text(s, T.MARGIN_IN + 0.20, y, T.col_w(4), 0.24, heading,
               size=T.SZ_MICRO, color=accent, bold=True, space=1.2)
        yy = y + 0.30
        for name, why in rows:
            T.text(s, T.MARGIN_IN + 0.20, yy, T.col_w(2), 0.28, name,
                   size=T.SZ_BODY, color=T.INK, bold=True, font=T.MONO)
            T.text(s, T.col_x(3), yy - 0.01, T.col_w(9), 0.38, why,
                   size=T.SZ_SMALL, color=T.MUTED, spacing=1.16)
            yy += 0.40
        y = yy + 0.22

    _footnote(s, "Recommendation: build on causal-learn and tigramite now; keep "
                 "GPUCSL as the scale path, since it needs no change of language.")
    return s


# ---------------------------------------------------------------------------
# 4 — Q3  time series and lag
# ---------------------------------------------------------------------------

def slide_04_timeseries(prs):
    s = T.blank(prs)
    y = _question(
        s, 4,
        "How do we handle time series and lag?",
        "Correct the test for autocorrelation, and take lags from the plant, "
        "not the data.",
        accent=T.AMBER)

    # left: the measurement
    T.text(s, T.MARGIN_IN, y, T.col_w(5), 0.26, "WHAT WE MEASURED",
           size=T.SZ_MICRO, color=T.AMBER, bold=True, space=1.2)
    yy = y + 0.34
    xs = [T.MARGIN_IN, T.col_x(3), T.col_x(4), T.col_x(5)]
    ws = [T.col_w(3), T.col_w(1), T.col_w(1), T.col_w(1)]
    yy = _colhead(s, yy, ["DATASET", "SAMPLES", "EFFECTIVE", "USABLE"], ws, xs)
    for name, n, ess, frac, col in [
            ("Ultra-processed food", "23,132", "121", "0.5 %", T.AMBER),
            ("Tennessee Eastman", "1,499", "530", "35 %", T.MUTED)]:
        T.text(s, xs[0], yy, ws[0], 0.3, name, size=T.SZ_BODY, color=T.INK)
        for i, v in enumerate([n, ess, frac], start=1):
            T.text(s, xs[i], yy, ws[i], 0.3, v, size=T.SZ_BODY, font=T.MONO,
                   color=col if i > 1 else T.INK,
                   bold=(i > 1))
        T.line(s, T.MARGIN_IN, yy + 0.40, T.col_x(6) - 0.1, yy + 0.40,
               color=T.RULE, pt=0.5)
        yy += 0.52

    T.text(s, T.MARGIN_IN, yy + 0.14, T.col_w(5), 0.9,
           [[("23,132 samples carry about 120 independent observations.",
              {"bold": True, "color": T.INK})],
            [("The smallest detectable correlation is 0.44, not 0.03 — so "
              "76 % of the edges we report are below what the evidence supports.",
              {"color": T.MUTED})]],
           size=T.SZ_SMALL, spacing=1.22)

    # right: what we do
    bx = T.col_x(6) + 0.20
    T.text(s, bx, y, T.col_w(6), 0.26, "WHAT WE DO ABOUT IT",
           size=T.SZ_MICRO, color=T.TEAL, bold=True, space=1.2)
    steps = [
        ("Correct the test",
         "Use the effective sample size in the Fisher-z degrees of freedom, and "
         "report both calibrations side by side."),
        ("Move to PCMCI+",
         "Built for autocorrelated series; “benefits from autocorrelation” "
         "rather than being broken by it. LPCMCI for UF, where the data authors "
         "record hidden confounders."),
        ("Set τₘₐₓ from plant knowledge",
         "We measured that the lags are not recoverable from the data: the series "
         "decorrelate over ~346 samples, seven times the longest real delay."),
    ]
    yy = y + 0.34
    for i, (head, body) in enumerate(steps, start=1):
        T.text(s, bx, yy, 0.3, 0.26, "%d" % i, size=T.SZ_SMALL, color=T.TEAL,
               bold=True, font=T.MONO)
        T.text(s, bx + 0.34, yy - 0.02, T.col_w(6) - 0.34, 0.28, head,
               size=T.SZ_BODY, color=T.INK, bold=True)
        T.text(s, bx + 0.34, yy + 0.26, T.col_w(6) - 0.34, 0.72, body,
               size=T.SZ_SMALL, color=T.MUTED, spacing=1.18)
        yy += 1.02

    _footnote(s, "Thinning the series was considered and rejected: it costs recall "
                 "heavily while barely moving precision, because it discards data "
                 "instead of correcting the test.")
    return s


# ---------------------------------------------------------------------------
# 5 — Q4  comparing graphs without ground truth
# ---------------------------------------------------------------------------

METRICS = [
    ("Edge stability",
     "% of bootstrap resamples containing the edge",
     "Gives every edge a confidence. Unstable edges are finite-sample noise.",
     T.TEAL),
    ("Self-compatibility",
     "disagreement across variable subsets",
     "A true edge should survive dropping other variables. Falsifies only.",
     T.TEAL),
    ("Falsification p-value",
     "violations vs a node-permutation baseline",
     "Answers “how many errors are too many?” — is this graph better "
     "than random?",
     T.SLATE),
    ("Held-out log-likelihood + BIC",
     "fit on a time-blocked split",
     "The only one that can rank OUR graph against the assumed ground truth.",
     T.AMBER),
]


def slide_05_metrics(prs):
    s = T.blank(prs)
    y = _question(
        s, 5,
        "Without ground truth, which metric?",
        "Four, in order of what each can rule out — no single number does it.",
        accent=T.SLATE)

    xs = [T.MARGIN_IN, T.col_x(3) + 0.1, T.col_x(6) + 0.2]
    ws = [T.col_w(3), T.col_w(3), T.col_w(6) - 0.2]
    y = _colhead(s, y, ["METRIC", "WHAT IT OUTPUTS", "WHAT IT BUYS US"], ws, xs)
    for name, out, buys, accent in METRICS:
        T.rect(s, T.MARGIN_IN - 0.16, y - 0.04, 0.035, 0.34, fill=accent)
        T.text(s, xs[0], y, ws[0], 0.3, name, size=T.SZ_BODY, color=T.INK, bold=True)
        T.text(s, xs[1], y, ws[1], 0.34, out, size=T.SZ_SMALL, color=T.MUTED,
               font=T.MONO, spacing=1.12)
        T.text(s, xs[2], y, ws[2], 0.34, buys, size=T.SZ_SMALL, color=T.INK,
               spacing=1.12)
        T.line(s, T.MARGIN_IN, y + 0.44, T.SLIDE_W_IN - T.MARGIN_IN, y + 0.44,
               color=T.RULE, pt=0.5)
        y += 0.60

    y += 0.16
    T.rect(s, T.MARGIN_IN, y, T.col_w(12), 0.90, fill=T.AMBER_SOFT, outline=None,
           radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.rect(s, T.MARGIN_IN, y, 0.045, 0.90, fill=T.AMBER)
    T.text(s, T.MARGIN_IN + 0.24, y + 0.13, T.col_w(12) - 0.5, 0.26,
           "The limit to state before anyone proposes BIC",
           size=T.SZ_BODY, color=T.INK, bold=True)
    T.text(s, T.MARGIN_IN + 0.24, y + 0.44, T.col_w(12) - 0.5, 0.40,
           "No observational score can separate graphs inside one Markov "
           "equivalence class — they share a likelihood. That is a theorem, not "
           "a tuning problem. It still works ACROSS classes, which is where our "
           "graph and the assumed ground truth sit.",
           size=T.SZ_SMALL, color=T.MUTED, spacing=1.18)
    return s


# ---------------------------------------------------------------------------
# 6 — Q5  error attribution
# ---------------------------------------------------------------------------

def slide_06_attribution(prs):
    s = T.blank(prs)
    y = _question(
        s, 6,
        "A score is bad — how do we find out why?",
        "Compute the ceiling first. What is left is ours to fix.",
        accent=T.TEAL)

    # the ladder of diagnostics
    T.text(s, T.MARGIN_IN, y, T.col_w(5), 0.26, "FOUR DIAGNOSTICS, IN ORDER",
           size=T.SZ_MICRO, color=T.TEAL, bold=True, space=1.2)
    yy = y + 0.34
    for head, body in [
            ("Identifiability ceiling",
             "Run the algorithm against the true graph's own independences. What it "
             "still cannot recover, no amount of data ever will."),
            ("Effective sample size",
             "How much independent evidence the samples actually carry."),
            ("Edges below detection",
             "Share of reported edges too weak for an honest test to see."),
            ("Bootstrap stability",
             "Stable but wrong ⇒ assumption failure. Unstable ⇒ too little data.")]:
        T.text(s, T.MARGIN_IN, yy, T.col_w(5), 0.26, head,
               size=T.SZ_BODY, color=T.INK, bold=True)
        T.text(s, T.MARGIN_IN, yy + 0.26, T.col_w(5), 0.56, body,
               size=T.SZ_SMALL, color=T.MUTED, spacing=1.16)
        yy += 0.88

    # the result: ceiling vs achieved
    bx = T.col_x(6) + 0.20
    bw = T.col_w(6)
    T.text(s, bx, y, bw, 0.26, "WHAT IT TOLD US", size=T.SZ_MICRO,
           color=T.SLATE, bold=True, space=1.2)
    yy = y + 0.40
    for name, ceiling, achieved in [("Ultra-processed food", 0.916, 0.036),
                                    ("Tennessee Eastman", 0.786, 0.143)]:
        T.text(s, bx, yy, bw, 0.26, name, size=T.SZ_BODY, color=T.INK, bold=True)
        track = bw - 1.05
        T.rect(s, bx, yy + 0.34, track, 0.20, fill=T.PAPER_DIM, outline=None)
        T.rect(s, bx, yy + 0.34, track * ceiling, 0.20, fill=T.SLATE_SOFT, outline=None)
        T.rect(s, bx, yy + 0.34, track * achieved, 0.20, fill=T.TEAL, outline=None)
        T.text(s, bx + track + 0.10, yy + 0.30, 0.95, 0.26,
               "%.0f%% of %.0f%%" % (achieved * 100, ceiling * 100),
               size=T.SZ_SMALL, color=T.MUTED, font=T.MONO)
        yy += 0.86

    T.legend(s, bx, yy - 0.10,
             [(T.TEAL, "what we achieve", None),
              (T.SLATE_SOFT, "what is achievable", None)], gap=1.85)

    T.text(s, bx, yy + 0.42, bw, 0.86,
           [[("The two datasets fail for different reasons.", {"bold": True})],
            [("UF: the ceiling is 92 % and we reach 4 % — that gap is our "
              "orientation stage, not the data. TE: it recovers the "
              "instrumentation, not the chemistry — 82 % of its correct edges "
              "are a valve and its own flow meter.", {"color": T.MUTED})]],
           size=T.SZ_SMALL, spacing=1.18)
    return s


# ---------------------------------------------------------------------------
# 7 — closing ask
# ---------------------------------------------------------------------------

def slide_07_ask(prs):
    s = T.blank(prs, dark=True)
    T.header(s, ACT, 7, accent=T.TEAL, dark=True)
    T.text(s, T.MARGIN_IN, T.MARGIN_IN + 0.70, T.col_w(9), 0.6,
           "Two decisions we cannot make ourselves", size=T.SZ_TITLE,
           color=T.ON_DARK, bold=True)

    asks = [
        ("Edge direction on control loops",
         "A valve drives its flow; the controller drives the valve from the "
         "measurement. Both are real — the pair is a feedback loop. The benchmark "
         "picks one. Do we follow its convention, or model the loop honestly and "
         "accept that it cannot be scored?"),
        ("A cyclic method for UFIMD",
         "That ground truth has 25 two-cycles, so an acyclic algorithm has no valid "
         "target. We found no maintained cyclic implementation of the quality we "
         "need. Do you know of one?"),
    ]
    y = 2.30
    for i, (head, body) in enumerate(asks, start=1):
        T.rect(s, T.MARGIN_IN, y, 0.045, 1.30, fill=T.TEAL)
        T.text(s, T.MARGIN_IN + 0.26, y - 0.02, 0.4, 0.26, "0%d" % i,
               size=T.SZ_MICRO, color=T.TEAL, bold=True, font=T.MONO, space=1.2)
        T.text(s, T.MARGIN_IN + 0.26, y + 0.26, T.col_w(9), 0.30, head,
               size=T.SZ_SUB, color=T.ON_DARK, bold=True)
        T.text(s, T.MARGIN_IN + 0.26, y + 0.62, T.col_w(9), 0.70, body,
               size=T.SZ_SMALL, color=T.RGBColor(0x9A, 0xA2, 0xAA), spacing=1.22)
        y += 1.72

    T.line(s, T.MARGIN_IN, 6.30, T.SLIDE_W_IN - T.MARGIN_IN, 6.30,
           color=T.RGBColor(0x33, 0x39, 0x40), pt=0.75)
    T.text(s, T.MARGIN_IN, 6.46, T.col_w(12), 0.3,
           "Full argument, alternatives and citations: FEEDBACK_RESPONSE.md",
           size=T.SZ_CAPTION, color=T.RGBColor(0x7A, 0x80, 0x88))
    return s


SLIDES = [
    slide_01_title,
    slide_02_cpp,
    slide_03_libraries,
    slide_04_timeseries,
    slide_05_metrics,
    slide_06_attribution,
    slide_07_ask,
]
