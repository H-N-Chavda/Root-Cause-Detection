"""Geometry check for the built deck.

There is no renderer on this machine, so overflow is caught arithmetically:
estimate how many lines each textbox actually needs at its width and point
size, compare against the box, and report anything that spills past its own
height or past the canvas. Conservative on purpose -- a warning here is a
thing to look at, not necessarily a defect.
"""

import sys
from math import ceil
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.65
# Average glyph advance as a fraction of point size, for a humanist sans.
GLYPH = 0.50
LINE = 1.22


def inches(value):
    return Emu(value).inches


def needed_height(shape):
    tf = shape.text_frame
    w = inches(shape.width)
    total = 0.0
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        size = max([r.font.size.pt for r in p.runs if r.font.size] or [12])
        cpl = max(int(w * 72 / (GLYPH * size)), 1)
        lines = max(ceil(len(text) / cpl), 1) if text else 1
        total += lines * size * LINE / 72
    return total


def main(path):
    prs = Presentation(path)
    problems = 0
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            x, y = inches(shape.left), inches(shape.top)
            w, h = inches(shape.width), inches(shape.height)
            tag = "slide %2d" % n
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                print("%s  OFF-CANVAS   %-46s  x=%.2f y=%.2f w=%.2f h=%.2f"
                      % (tag, repr(shape.name)[:44], x, y, w, h))
                problems += 1
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            need = needed_height(shape)
            bottom = y + need
            if need > h + 0.06:
                print("%s  OVERFLOWS    %-46s  needs %.2f in, box %.2f in, "
                      "ends %.2f" % (tag, text.replace("\n", " / ")[:44], need, h, bottom))
                problems += 1
            elif bottom > SLIDE_H - 0.10:
                print("%s  PAST FOOT    %-46s  ends %.2f"
                      % (tag, text.replace("\n", " / ")[:44], bottom))
                problems += 1
        problems += _overlaps(slide, n)
    print("\n%d geometry warnings" % problems)
    return problems


def _overlaps(slide, n):
    """Text blocks that would print on top of each other.

    Uses the estimated rendered height rather than the declared box height,
    which is what actually collides once PowerPoint lays the text out.
    """
    boxes = []
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        x, y = inches(shape.left), inches(shape.top)
        boxes.append((x, y, x + inches(shape.width), y + needed_height(shape),
                      shape.text_frame.text.replace("\n", " / ")))
    hits = 0
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[2], b[2]) - max(a[0], b[0])
            oy = min(a[3], b[3]) - max(a[1], b[1])
            if ox > 0.12 and oy > 0.10:
                print("slide %2d  COLLIDES     %-30s  with  %-30s  (%.2f x %.2f in)"
                      % (n, a[4][:28], b[4][:28], ox, oy))
                hits += 1
    return hits


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1] if len(sys.argv) > 1 else
                       Path(__file__).parent / "Week4_PCMCIplus_VARLiNGAM_LSTE.pptx") else 0)
